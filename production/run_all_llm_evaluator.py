import time
import importlib
import sys
import os
import pandas as pd
import openpyxl
from typing import List, Dict, Any, Tuple

from faithfullness_eval import FaithfulnessEvaluator
from response_relevancy_eval import ResponseRelevancyEvaluator
from config.config import Config
from config.logger import Logger
from utils.constants import PROJECT_ROOT, EMBEDDING_DEPARTMENT_NAME


class LLMEvaluationRunner:
    """
    Comprehensive LLM evaluation runner that tests multiple LLM models
    across different retrieval methods and evaluation metrics.
    """
    
    def __init__(self):
        self.config = Config()
        self.logger = Logger(log_file="llm_evaluation_runner.log").get_logger()
        
        # LLM models to evaluate
        self.llm_models = [
        "llama3-8b-8192",
        "llama-3.3-70b-versatile", 
        "llama-3.1-8b-instant",
        # "llama3-70b-8192"  # DEPRECATED - model decommissioned
        ]
        
        # Retrieval methods to test
        self.retrieval_methods = [
            "faiss_reranker", 
            "chroma_metadata", 
            "chroma_metadata_reranker"
        ]
        
        # Evaluation metrics and their corresponding classes
        self.evaluators = {
            "faithfulness": (FaithfulnessEvaluator, "run_faithfulness"),
            "response_relevancy": (ResponseRelevancyEvaluator, "run_response_relevancy")
        }
        
        # Set up output paths
        self.embedding_department_name = EMBEDDING_DEPARTMENT_NAME
        self.output_path = os.path.join(PROJECT_ROOT, "data", self.embedding_department_name, "llm_evaluation_results.xlsx")
        
        self.logger.info("LLM Evaluation Runner initialized")
        self.logger.info(f"Models to evaluate: {self.llm_models}")
        self.logger.info(f"Retrieval methods: {self.retrieval_methods}")
        self.logger.info(f"Evaluation metrics: {list(self.evaluators.keys())}")
        self.logger.info(f"Results will be saved to: {self.output_path}")

    def update_llm_model(self, model_name: str):
        """Update the configuration to use a specific LLM model"""
        self.config.groq_model_name = model_name
        self.logger.info(f"Updated LLM model to: {model_name}")

    def get_sheet_name(self, eval_name: str, retrieval_method: str, llm_model: str) -> str:
        """Generate a unique sheet name for each combination"""
        # Clean model name for sheet name (Excel sheet names have restrictions)
        clean_model = llm_model.replace("-", "_").replace(".", "_")
        
        # Map evaluation names to shorter versions for sheet names
        eval_short_names = {

            "faithfulness": "faith",
            "response_relevancy": "resp_rel"
        }
        
        # Map retrieval method names to shorter versions
        retrieval_short_names = {
            "faiss_reranker": "faiss",
            "chroma_metadata": "chroma",
            "chroma_metadata_reranker": "chroma_rr"
        }
        
        eval_short = eval_short_names.get(eval_name, eval_name)
        retrieval_short = retrieval_short_names.get(retrieval_method, retrieval_method)
        
        return f"{eval_short}_{retrieval_short}_{clean_model}"

    def save_results_with_llm_info(self, results_df: pd.DataFrame, sheet_name: str, 
                                 llm_model: str, eval_name: str, retrieval_method: str):
        """Save results to Excel with LLM model information included"""
        try:
            # Add LLM model column to results
            results_df_copy = results_df.copy()
            results_df_copy['LLM_Model'] = llm_model
            results_df_copy['Evaluation_Metric'] = eval_name
            
            # Reorder columns to put LLM_Model and Evaluation_Metric first
            cols = ['LLM_Model', 'Evaluation_Metric'] + [col for col in results_df_copy.columns 
                                                        if col not in ['LLM_Model', 'Evaluation_Metric']]
            results_df_copy = results_df_copy[cols]
            
            # Save to Excel
            if os.path.exists(self.output_path):
                try:
                    book = openpyxl.load_workbook(self.output_path)
                    if sheet_name in book.sheetnames:
                        book.remove(book[sheet_name])
                        book.save(self.output_path)
                except Exception as e:
                    self.logger.warning(f"Error modifying existing workbook: {e}")

            with pd.ExcelWriter(self.output_path, engine='openpyxl', 
                              mode='a' if os.path.exists(self.output_path) else 'w') as writer:
                results_df_copy.to_excel(writer, sheet_name=sheet_name, index=False)

            self.logger.info(f"Results saved to sheet '{sheet_name}' in {self.output_path}")
            
        except Exception as e:
            self.logger.error(f"Error saving results to Excel: {e}")
            raise

    def run_single_evaluation(self, evaluator_class, method_name: str, 
                            retrieval_method: str, llm_model: str) -> Dict[str, Any]:
        """
        Run a single evaluation for a specific evaluator, retrieval method, and LLM model.
        """
        try:
            # Update LLM model in config
            self.update_llm_model(llm_model)
            
            # Create evaluator instance with updated config
            evaluator_instance = evaluator_class(self.config)
            
            # Get the method to call
            run_method = getattr(evaluator_instance, method_name)
            
            # Run the evaluation
            self.logger.info(f"Running {method_name} with {retrieval_method} using {llm_model}")
            results = run_method(retrieval_method=retrieval_method)
            
            # Calculate average score
            metric_column = results.columns[4]  # The metric column is typically at index 4
            avg_score = results[metric_column].mean()
            
            # Generate sheet name and save results with LLM info
            sheet_name = self.get_sheet_name(method_name.replace("run_", ""), retrieval_method, llm_model)
            self.save_results_with_llm_info(results, sheet_name, llm_model, 
                                          method_name.replace("run_", ""), retrieval_method)
            
            result_info = {
                "llm_model": llm_model,
                "retrieval_method": retrieval_method,
                "evaluation_metric": method_name,
                "average_score": avg_score,
                "total_samples": len(results),
                "successful_evaluations": results[metric_column].notna().sum(),
                "status": "success",
                "sheet_name": sheet_name
            }
            
            self.logger.info(f"✅ Completed {method_name} for {retrieval_method} using {llm_model}")
            self.logger.info(f"Average score: {avg_score}")
            
            return result_info
            
        except Exception as e:
            error_info = {
                "llm_model": llm_model,
                "retrieval_method": retrieval_method,
                "evaluation_metric": method_name,
                "average_score": None,
                "total_samples": 0,
                "successful_evaluations": 0,
                "status": "error",
                "error_message": str(e),
                "sheet_name": None
            }
            
            self.logger.error(f"❌ Error in {method_name} with {retrieval_method} using {llm_model}: {str(e)}")
            return error_info

    def create_master_summary_sheet(self, all_results: Dict[str, Dict[str, List[Dict[str, Any]]]]):
        """Create a comprehensive summary sheet with all results"""
        summary_data = []
        
        for eval_name, retrieval_results in all_results.items():
            for retrieval_method, model_results in retrieval_results.items():
                for result in model_results:
                    if result['status'] == 'success':
                        summary_data.append({
                            'Evaluation_Metric': eval_name,
                            'Retrieval_Method': retrieval_method,
                            'LLM_Model': result['llm_model'],
                            'Average_Score': result['average_score'],
                            'Total_Samples': result['total_samples'],
                            'Successful_Evaluations': result['successful_evaluations'],
                            'Success_Rate': result['successful_evaluations'] / result['total_samples'] * 100,
                            'Sheet_Name': result.get('sheet_name', 'N/A')
                        })
        
        if summary_data:
            summary_df = pd.DataFrame(summary_data)
            
            try:
                # Remove existing summary sheet if it exists
                if os.path.exists(self.output_path):
                    try:
                        book = openpyxl.load_workbook(self.output_path)
                        if 'Master_Summary' in book.sheetnames:
                            book.remove(book['Master_Summary'])
                            book.save(self.output_path)
                    except Exception as e:
                        self.logger.warning(f"Error removing existing summary: {e}")
                
                # Save summary sheet
                with pd.ExcelWriter(self.output_path, engine='openpyxl', 
                                  mode='a' if os.path.exists(self.output_path) else 'w') as writer:
                    summary_df.to_excel(writer, sheet_name='Master_Summary', index=False)
                
                self.logger.info("Master summary sheet created successfully")
                
            except Exception as e:
                self.logger.error(f"Error creating master summary sheet: {e}")

    def create_best_model_analysis(self, all_results: Dict[str, Dict[str, List[Dict[str, Any]]]]):
        """Create analysis showing best performing models for each metric/retrieval combination"""
        analysis_data = []
        
        for eval_name, retrieval_results in all_results.items():
            for retrieval_method, model_results in retrieval_results.items():
                # Find best performing model for this combination
                successful_results = [r for r in model_results if r['status'] == 'success' and r['average_score'] is not None]
                
                if successful_results:
                    best_result = max(successful_results, key=lambda x: x['average_score'])
                    worst_result = min(successful_results, key=lambda x: x['average_score'])
                    
                    analysis_data.append({
                        'Evaluation_Metric': eval_name,
                        'Retrieval_Method': retrieval_method,
                        'Best_Model': best_result['llm_model'],
                        'Best_Score': best_result['average_score'],
                        'Worst_Model': worst_result['llm_model'],
                        'Worst_Score': worst_result['average_score'],
                        'Score_Difference': best_result['average_score'] - worst_result['average_score'],
                        'Models_Tested': len(successful_results)
                    })
        
        if analysis_data:
            analysis_df = pd.DataFrame(analysis_data)
            
            try:
                # Remove existing analysis sheet if it exists
                if os.path.exists(self.output_path):
                    try:
                        book = openpyxl.load_workbook(self.output_path)
                        if 'Best_Model_Analysis' in book.sheetnames:
                            book.remove(book['Best_Model_Analysis'])
                            book.save(self.output_path)
                    except Exception as e:
                        self.logger.warning(f"Error removing existing analysis: {e}")
                
                # Save analysis sheet
                with pd.ExcelWriter(self.output_path, engine='openpyxl', 
                                  mode='a' if os.path.exists(self.output_path) else 'w') as writer:
                    analysis_df.to_excel(writer, sheet_name='Best_Model_Analysis', index=False)
                
                self.logger.info("Best model analysis sheet created successfully")
                
            except Exception as e:
                self.logger.error(f"Error creating best model analysis sheet: {e}")

    def run_evaluator_for_all_models(self, evaluator_name: str, evaluator_class, 
                                   method_name: str, retrieval_method: str) -> List[Dict[str, Any]]:
        """Run a specific evaluator across all LLM models for a given retrieval method."""
        results = []
        
        self.logger.info(f"\n{'='*100}")
        self.logger.info(f"Running {evaluator_name} with {retrieval_method}")
        self.logger.info(f"{'='*100}")
        
        for i, llm_model in enumerate(self.llm_models):
            self.logger.info(f"\n{'-'*80}")
            self.logger.info(f"Model {i+1}/{len(self.llm_models)}: {llm_model}")
            self.logger.info(f"{'-'*80}")
            
            result = self.run_single_evaluation(
                evaluator_class, method_name, retrieval_method, llm_model
            )
            results.append(result)
            
            # Add delay between models to prevent rate limiting
            if i < len(self.llm_models) - 1:
                self.logger.info("Waiting 30 seconds before next model...")
                time.sleep(30)
        
        return results

    def run_all_evaluations(self) -> Dict[str, Dict[str, List[Dict[str, Any]]]]:
        """Run all evaluations across all combinations of models, retrieval methods, and metrics."""
        all_results = {}
        total_combinations = (len(self.evaluators) * len(self.retrieval_methods) * 
                            len(self.llm_models))
        
        self.logger.info(f"\n Starting comprehensive LLM evaluation")
        self.logger.info(f"Total combinations to evaluate: {total_combinations}")
        self.logger.info(f"Estimated time: ~{total_combinations * 2} minutes")
        self.logger.info(f"Results will be saved to: {self.output_path}")
        
        evaluation_counter = 0
        
        # Iterate through each evaluator
        for eval_name, (evaluator_class, method_name) in self.evaluators.items():
            all_results[eval_name] = {}
            
            # Iterate through each retrieval method
            for retrieval_method in self.retrieval_methods:
                self.logger.info(f"\n Progress: {evaluation_counter}/{len(self.evaluators) * len(self.retrieval_methods)} evaluator-method combinations")
                
                # Run evaluator for all models with this retrieval method
                results = self.run_evaluator_for_all_models(
                    eval_name, evaluator_class, method_name, retrieval_method
                )
                
                all_results[eval_name][retrieval_method] = results
                evaluation_counter += 1
                
                # Add delay between retrieval methods
                if retrieval_method != self.retrieval_methods[-1]:
                    self.logger.info("Waiting 60 seconds before next retrieval method...")
                    time.sleep(60)
            
            # Add delay between evaluators
            if eval_name != list(self.evaluators.keys())[-1]:
                self.logger.info("Waiting 120 seconds before next evaluator...")
                time.sleep(120)
        
        # Create summary sheets
        self.create_master_summary_sheet(all_results)
        self.create_best_model_analysis(all_results)
        
        return all_results

    def generate_summary_report(self, all_results: Dict[str, Dict[str, List[Dict[str, Any]]]]):
        """Generate a comprehensive summary report of all evaluations"""
        
        self.logger.info(f"\n{'='*100}")
        self.logger.info("COMPREHENSIVE LLM EVALUATION SUMMARY")
        self.logger.info(f"{'='*100}")
        self.logger.info(f"📁 Results saved to: {self.output_path}")
        
        # Summary by LLM model
        model_summary = {}
        for llm_model in self.llm_models:
            model_summary[llm_model] = {
                "total_evaluations": 0,
                "successful_evaluations": 0,
                "average_scores": {},
                "best_combinations": [],
                "worst_combinations": []
            }
        
        # Process results
        for eval_name, retrieval_results in all_results.items():
            self.logger.info(f"\n{'-'*80}")
            self.logger.info(f"📊 {eval_name.upper()} RESULTS")
            self.logger.info(f"{'-'*80}")
            
            for retrieval_method, model_results in retrieval_results.items():
                self.logger.info(f"\n🔍 Retrieval Method: {retrieval_method}")
                
                best_score = -1
                worst_score = float('inf')
                best_model = ""
                worst_model = ""
                
                for result in model_results:
                    model = result['llm_model']
                    score = result['average_score']
                    status = result['status']
                    sheet_name = result.get('sheet_name', 'N/A')
                    
                    # Update model summary
                    model_summary[model]["total_evaluations"] += 1
                    if status == "success":
                        model_summary[model]["successful_evaluations"] += 1
                        
                        if eval_name not in model_summary[model]["average_scores"]:
                            model_summary[model]["average_scores"][eval_name] = {}
                        model_summary[model]["average_scores"][eval_name][retrieval_method] = score
                        
                        # Track best and worst for this retrieval method
                        if score is not None:
                            if score > best_score:
                                best_score = score
                                best_model = model
                            if score < worst_score:
                                worst_score = score
                                worst_model = model
                    
                    # Log individual result
                    status_icon = "✅" if status == "success" else "❌"
                    score_str = f"{score:.4f}" if score is not None else "Failed"
                    self.logger.info(f"  {status_icon} {model}: {score_str} (Sheet: {sheet_name})")
                
                # Log best and worst for this combination
                if best_model and worst_model:
                    self.logger.info(f"  🏆 Best: {best_model} ({best_score:.4f})")
                    self.logger.info(f"  📉 Worst: {worst_model} ({worst_score:.4f})")
        
        # Overall model performance summary
        self.logger.info(f"\n{'='*100}")
        self.logger.info("🏆 OVERALL MODEL PERFORMANCE SUMMARY")
        self.logger.info(f"{'='*100}")
        
        for model, summary in model_summary.items():
            success_rate = (summary["successful_evaluations"] / 
                          summary["total_evaluations"] * 100 if summary["total_evaluations"] > 0 else 0)
            
            self.logger.info(f"\n📈 {model}:")
            self.logger.info(f"  Success Rate: {success_rate:.1f}% ({summary['successful_evaluations']}/{summary['total_evaluations']})")
            
            if summary["average_scores"]:
                self.logger.info("  Average Scores by Metric:")
                for eval_name, retrieval_scores in summary["average_scores"].items():
                    avg_across_retrievals = sum(retrieval_scores.values()) / len(retrieval_scores)
                    best_retrieval = max(retrieval_scores.items(), key=lambda x: x[1])
                    self.logger.info(f"    {eval_name}: {avg_across_retrievals:.4f} (best: {best_retrieval[0]} - {best_retrieval[1]:.4f})")

        self.logger.info(f"\n📊 EXCEL FILE CONTENTS:")
        self.logger.info(f"  📁 File Location: {self.output_path}")
        self.logger.info(f"  📋 Individual result sheets: One per LLM-retrieval-metric combination")
        self.logger.info(f"  📋 Master_Summary sheet: All results in one table")
        self.logger.info(f"  📋 Best_Model_Analysis sheet: Best/worst models per combination")
        
        self.logger.info(f"\n🎉 Evaluation completed successfully!")


def main():
    """Main execution function"""
    print("🚀 Starting LLM Evaluation Suite")
    print("="*80)
    
    try:
        # Initialize the evaluation runner
        runner = LLMEvaluationRunner()
        
        # Run all evaluations
        all_results = runner.run_all_evaluations()
        
        # Generate summary report
        runner.generate_summary_report(all_results)
        
        print("\n✅ All LLM evaluations completed successfully!")
        print(f"📊 Check the Excel file at: {runner.output_path}")
        print("📋 The file contains individual sheets for each combination plus summary sheets.")
        
    except KeyboardInterrupt:
        print("\n⚠️  Evaluation interrupted by user")
    except Exception as e:
        print(f"\n❌ Error during evaluation: {str(e)}")
        raise


if __name__ == "__main__":
    main()